"""Dựng deck thuyết trình Đề tài 8 trên template PowerPoint của IUH.

Học viên: Phan Nguyễn Mai Phương · MSHV 25002711.

Deck gồm 20 slide: giới thiệu, dữ liệu và hạ tầng, phân tích và dự đoán, kết luận.
Mọi số liệu đọc từ artifact đã nghiệm thu; script không huấn luyện lại mô hình.

Chạy từ thư mục gốc repository:

    uv run --frozen python slides/extract_slide_data.py
    uv run --frozen python -m cmapss.degradation
    uv run --frozen python slides/slide_figures.py
    uv run --no-project --with python-pptx==1.0.2 --with pillow==12.3.0 python slides/build_presentation.py
"""

from __future__ import annotations

import io
import json
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from cmapss.presentation import COLORS, FONT, format_number_vi as vn

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "slides/template/Template Powerpoint v2 (ty le 16-9).pptx"
OUT = ROOT / "slides/cmapss_fd001_presentation.pptx"
TOTAL = 20

D = json.loads((ROOT / "slides/slide_data.json").read_text(encoding="utf-8"))
DEG = D["degradation"]
T40 = D["maintenance"]["threshold_40_test"]
NOISE = {row["sensor"]: row for row in D["noise_summary"]}

# Màu lấy từ template IUH: xanh thương hiệu và vàng của đường viền footer.
BLUE, GOLD = COLORS["blue"].lstrip("#"), COLORS["gold"].lstrip("#")
INK, SLATE, LINE = "1B2733", "5A6B7B", "C9D2E3"
TINT, GOLD_TINT, WHITE = "EAF0FA", "FFF4DA", "FFFFFF"
RED, TEAL = COLORS["red"].lstrip("#"), COLORS["teal"].lstrip("#")
PARTS = [
    "Giới thiệu",
    "Dữ liệu và hạ tầng",
    "Phân tích và dự đoán",
    "Kết luận",
]

M, P, A = D["modeling"], D["pipeline"], D["alignment"]
FD = {row["name"]: row for row in D["subsets"]}
WARN = {f"{r['model']}_{r['threshold']}": r for r in A["warning_policies"]}
AN, AN_RAW, ACFG = DEG["anomaly_smoothed"], DEG["anomaly_raw_for_comparison"], DEG["anomaly_config"]
MNT = D["maintenance"]
HDFS = D["hdfs"]
assert HDFS["mapreduce_vs_pandas"]["exact_match"] and HDFS["spark_vs_pandas"]["exact_match"]
# Đối chiếu rolling features Pandas–PySpark ở cả ba lượt benchmark.
BENCH_PARITY = D["benchmark_parity"]
ROLLING_ROWS = {run["features"]["rows"] for run in BENCH_PARITY.values()}
assert len(ROLLING_ROWS) == 1 and all(run["features"]["all_cells_match"] for run in BENCH_PARITY.values())
SENS = D["sensitivity"]
# Sensor thay đổi cùng chiều ở mọi động cơ train (kiểm định đầu/cuối vòng đời).
SAME_DIRECTION = D["statistics"]["same_direction"]
assert {"sensor_4", "sensor_7", "sensor_11", "sensor_12", "sensor_20"} <= set(SAME_DIRECTION)
# Lời nói slide 15 khẳng định mọi cấu hình trong lưới đều phát hiện đủ động cơ.
assert all(row["n_detected"] == row["n_engines"] for row in SENS)


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


# ----------------------------------------------------------------------------- primitives
def text(slide, x, y, w, h, content, *, size=14, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, space_after=0, font=FONT):
    """Thêm textbox. ``content`` là chuỗi, hoặc danh sách đoạn; mỗi đoạn là chuỗi
    hoặc danh sách run ``(chữ, {"bold":..., "color":..., "size":...})``."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.auto_size = None
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = anchor
    paragraphs = content if isinstance(content, list) else [content]
    for index, para in enumerate(paragraphs):
        p = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        runs = para if isinstance(para, list) else [(para, {})]
        for chunk, style in runs:
            run = p.add_run()
            run.text = chunk
            run.font.name = font
            run.font.size = Pt(style.get("size", size))
            run.font.bold = style.get("bold", bold)
            run.font.color.rgb = rgb(style.get("color", color))
    return box


def rect(slide, x, y, w, h, fill, *, line=None, rounded=True):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    if rounded:
        shape.adjustments[0] = 0.08
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(0.75)
    else:
        shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def circle(slide, x, y, d, label, fill=BLUE, color=WHITE, size=13):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.fill.background()
    shape.shadow.inherit = False
    text(slide, x, y, d, d, label, size=size, color=color, bold=True, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE)


def arrow(slide, x1, y1, x2, y2, color=SLATE, *, head=True, dash=False):
    connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    connector.line.color.rgb = rgb(color)
    connector.line.width = Pt(1.5)
    if dash:
        connector.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    if head:
        line = connector.line._get_or_add_ln()
        tail = line.makeelement(qn("a:tailEnd"), {"type": "triangle"})
        line.append(tail)


def path(slide, points, color=SLATE, *, dash=False):
    """Đường gấp khúc qua các điểm; chỉ đoạn cuối có đầu mũi tên."""
    for (x1, y1), (x2, y2) in zip(points, points[1:]):
        arrow(slide, x1, y1, x2, y2, color, head=(x2, y2) == points[-1], dash=dash)


def picture(slide, rel, x, y, w, h, *, crop=None, align="center"):
    """Đặt hình vừa khung (giữ tỉ lệ); trả về (x, y, w, h) thực tế."""
    image = Image.open(ROOT / rel)
    if crop:
        image = image.crop(crop)
    iw, ih = image.size
    scale = min(w / iw, h / ih)
    pw, ph = iw * scale, ih * scale
    px = x + (w - pw) / 2 if align == "center" else x
    py = y + (h - ph) / 2 if align == "center" else y
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    buffer.seek(0)
    slide.shapes.add_picture(buffer, Inches(px), Inches(py), Inches(pw), Inches(ph))
    return px, py, pw, ph


def table(slide, rows, x, y, col_widths, *, row_h=0.42, size=12, highlight=None, align_right_from=None):
    n_rows, n_cols = len(rows), len(rows[0])
    shape = slide.shapes.add_table(n_rows, n_cols, Inches(x), Inches(y), Inches(sum(col_widths)), Inches(row_h * n_rows))
    tbl = shape.table
    tbl.first_row = True
    for j, width in enumerate(col_widths):
        tbl.columns[j].width = Inches(width)
    for i in range(n_rows):
        tbl.rows[i].height = Inches(row_h)
        for j in range(n_cols):
            cell = tbl.cell(i, j)
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if i == 0:
                cell.fill.fore_color.rgb = rgb(BLUE)
            elif highlight is not None and i == highlight:
                cell.fill.fore_color.rgb = rgb(GOLD_TINT)
            else:
                cell.fill.fore_color.rgb = rgb(TINT if i % 2 == 0 else WHITE)
            frame = cell.text_frame
            frame.word_wrap = True
            p = frame.paragraphs[0]
            p.text = ""
            if align_right_from is not None and j >= align_right_from and i > 0:
                p.alignment = PP_ALIGN.RIGHT
            run = p.add_run()
            run.text = str(rows[i][j])
            run.font.name = FONT
            run.font.size = Pt(size)
            run.font.bold = i == 0 or (highlight is not None and i == highlight)
            run.font.color.rgb = rgb(WHITE if i == 0 else INK)
    return shape


def bar_chart(slide, x, y, w, h, categories, series, colors, *, title, number_format="0", max_value=None,
              horizontal=False, legend=True, decimals=0):
    """``decimals`` > 0: hiện số thập phân với dấu phẩy kiểu Việt Nam ở mọi locale.

    PowerPoint lấy dấu thập phân theo máy chiếu, nên giá trị được nhân 10^decimals và
    nhãn chèn dấu phẩy cố định trước các chữ số cuối; trục giá trị đang ẩn nên không lộ thang đo.
    """
    scale = 10 ** decimals
    if decimals:
        number_format = "0\\," + "0" * decimals
        series = [(name, [round(v * scale) for v in values]) for name, values in series]
        max_value = None if max_value is None else max_value * scale
    data = CategoryChartData()
    data.categories = categories
    for name, values in series:
        data.add_series(name, values)
    kind = XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED
    chart = slide.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), data).chart
    chart.font.name = FONT
    chart.font.size = Pt(12)
    chart.font.color.rgb = rgb(INK)
    chart.has_title = True
    chart.chart_title.text_frame.text = title
    title_run = chart.chart_title.text_frame.paragraphs[0].runs[0]
    title_run.font.size, title_run.font.bold, title_run.font.name = Pt(13), True, FONT
    title_run.font.color.rgb = rgb(INK)
    chart.has_legend = legend
    if legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(11)
    plot = chart.plots[0]
    plot.gap_width = 70
    plot.has_data_labels = True
    labels = plot.data_labels
    labels.number_format = number_format
    labels.number_format_is_linked = False
    labels.position = XL_LABEL_POSITION.OUTSIDE_END
    labels.font.size = Pt(12)
    for item, color in zip(plot.series, colors):
        item.format.fill.solid()
        item.format.fill.fore_color.rgb = rgb(color)
    value_axis = chart.value_axis
    value_axis.visible = False
    value_axis.has_major_gridlines = False
    value_axis.minimum_scale = 0
    if max_value is not None:
        value_axis.maximum_scale = max_value
    chart.category_axis.tick_labels.font.size = Pt(12)
    chart.category_axis.format.line.color.rgb = rgb(LINE)
    if horizontal:
        chart.category_axis.reverse_order = True
    return chart


def stat(slide, x, y, w, value, label, *, color=BLUE, size=28):
    text(slide, x, y, w, 0.55, value, size=size, bold=True, color=color, anchor=MSO_ANCHOR.BOTTOM)
    text(slide, x, y + 0.6, w, 0.5, label, size=12, color=SLATE)


def numbers_right(slide, items):
    """Cột số to bên phải một hình đứng; mỗi mục là (giá trị, nhãn, màu)."""
    for i, (value, label, color) in enumerate(items):
        y = 2.3 + i * 1.4
        text(slide, 8.9, y, 3.5, 0.7, value, size=34, bold=True, color=color, anchor=MSO_ANCHOR.BOTTOM)
        text(slide, 8.9, y + 0.74, 3.5, 0.6, label, size=14, color=SLATE)


def numbers_below(slide, items):
    """Hàng số to dưới một hình nằm ngang; mỗi mục là (giá trị, nhãn, màu)."""
    for i, (value, label, color) in enumerate(items):
        x = 0.92 + i * 3.9
        text(slide, x, 5.55, 3.7, 0.6, value, size=30, bold=True, color=color, anchor=MSO_ANCHOR.BOTTOM)
        text(slide, x, 6.18, 3.7, 0.5, label, size=13, color=SLATE)


def wide_figure(slide, rel, *, crop=None):
    """Hình nằm ngang do pipeline sinh ra, căn giữa phía trên hàng số."""
    picture(slide, rel, 0.92, 2.0, 11.5, 3.45, crop=crop)


def tall_figure(slide, rel):
    """Hình đứng do pipeline sinh ra, bên trái cột số."""
    picture(slide, rel, 0.92, 2.05, 7.6, 4.7, align="left")


# Hình thật do notebook sinh ra; run được notebook trỏ tới, giữ lại khi đóng gói.
NB01 = "outputs/notebook01/run-_uov4vhx/figures"


# ----------------------------------------------------------------------------- slide frames
PRS = Presentation(str(TEMPLATE))
# Xóa 10 slide mẫu, giữ nguyên master và layout của template.
for slide_id in list(PRS.slides._sldIdLst):
    PRS.part.drop_rel(slide_id.rId)
    PRS.slides._sldIdLst.remove(slide_id)
LAYOUT = {layout.name: layout for layout in PRS.slide_layouts}


def _remove(placeholder) -> None:
    element = placeholder._element
    element.getparent().remove(element)


def content(part: int, n: int, title: str, takeaway: str, source: str, notes: str):
    """Slide nội dung trên layout Title and Content của template."""
    slide = PRS.slides.add_slide(LAYOUT["Title and Content"])
    for placeholder in list(slide.placeholders):
        if placeholder.placeholder_format.idx != 0:
            _remove(placeholder)
    heading = slide.shapes.title
    heading.left, heading.top, heading.width, heading.height = Inches(0.92), Inches(1.3), Inches(11.5), Inches(0.62)
    heading.text_frame.word_wrap = True
    heading.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    heading.text = title
    run = heading.text_frame.paragraphs[0].runs[0]
    run.font.size, run.font.bold, run.font.name = Pt(26), True, FONT
    run.font.color.rgb = rgb(BLUE)
    text(slide, 6.2, 0.42, 6.2, 0.38, f"PHẦN {part} · {PARTS[part - 1].upper()}", size=12, bold=True,
         color=WHITE, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    if takeaway:
        text(slide, 0.92, 1.93, 11.5, 0.42, takeaway, size=15, color=SLATE, anchor=MSO_ANCHOR.MIDDLE)
    text(slide, 11.2, 6.98, 1.22, 0.32, f"{n} / {TOTAL}", size=11, color=SLATE, align=PP_ALIGN.RIGHT,
         anchor=MSO_ANCHOR.MIDDLE)
    # Nguồn số liệu để trong ghi chú người nói, không in lên slide.
    slide.notes_slide.notes_text_frame.text = f"{notes}\nNguồn: {source}"
    return slide


GOLD_DARK = "B37800"
DP = D["data_platform"]
CODE_BG = "F4F6FA"


def code_box(slide, x, y, w, h, lines, *, size=11):
    """Khung chữ đơn cách cho câu lệnh, truy vấn và log thật."""
    rect(slide, x, y, w, h, CODE_BG, line=LINE)
    text(slide, x + 0.15, y + 0.1, w - 0.3, h - 0.2, lines, size=size, font="Consolas", space_after=1)


def plain_list(slide, x, y, w, h, items, *, size=17):
    """Danh sách gạch đầu dòng bằng ký tự, không dùng định dạng bullet của template."""
    text(slide, x, y, w, h, [f"–  {item}" for item in items], size=size, space_after=10)


# ============================================================================ 1. Bìa
def slide_cover():
    slide = PRS.slides.add_slide(LAYOUT["Title Slide"])
    for placeholder in list(slide.placeholders):
        if placeholder.placeholder_format.idx not in (0, 1):
            _remove(placeholder)
    title, subtitle = slide.placeholders[0], slide.placeholders[1]
    title.left, title.top, title.width, title.height = Inches(0.92), Inches(1.75), Inches(11.5), Inches(1.5)
    title.text = "Phân tích dữ liệu IoT và bảo trì dự đoán"
    p = title.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.runs[0].font.size, p.runs[0].font.bold, p.runs[0].font.name = Pt(40), True, FONT
    p.runs[0].font.color.rgb = rgb(BLUE)
    subtitle.left, subtitle.top, subtitle.width, subtitle.height = Inches(1.4), Inches(3.25), Inches(10.5), Inches(1.0)
    # Gạch nối không ngắt (U+2011) để "C-MAPSS" không bị tách dòng.
    subtitle.text = "Phát hiện bất thường, dự đoán RUL và cảnh báo sớm trên dữ liệu cảm biến NASA C‑MAPSS"
    p = subtitle.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    for run in p.runs:
        run.font.size, run.font.name = Pt(18), FONT
        run.font.color.rgb = rgb(INK)
    text(slide, 0.92, 1.3, 11.5, 0.4, "BÁO CÁO CUỐI KỲ · PHÂN TÍCH DỮ LIỆU LỚN · ĐỀ TÀI 8", size=14, bold=True,
         color=GOLD_DARK, align=PP_ALIGN.CENTER)
    text(slide, 0.92, 4.75, 11.5, 1.6, [
        [("Giảng viên hướng dẫn: ", {"color": SLATE}), ("TS. Võ Thị Thanh Vân", {"bold": True, "color": BLUE})],
        [("Học viên thực hiện: ", {"color": SLATE}),
         ("Phan Nguyễn Mai Phương · MSHV 25002711", {"bold": True, "color": BLUE})],
        [("Ngành Khoa học máy tính · Tháng 10/2026", {"size": 14, "color": INK})],
    ], size=18, align=PP_ALIGN.CENTER, space_after=8)
    slide.notes_slide.notes_text_frame.text = (
        "Thời lượng: 20 giây.\nKính chào quý thầy cô. Em là Phan Nguyễn Mai Phương, MSHV 25002711. Em trình bày "
        "Đề tài 8: phân tích dữ liệu cảm biến IoT để phát hiện dấu hiệu bất thường, dự đoán thời điểm cần bảo trì và "
        "cảnh báo sớm, dưới sự hướng dẫn của TS. Võ Thị Thanh Vân.")


# ============================================================================ 2. Bài toán
def slide_problem():
    ev = MNT["evaluation"]
    failure = int(M["example_warning"]["failure_cycle"])
    s = content(1, 2, "Bài toán: xác định thời điểm cần bảo trì cho từng động cơ", "",
                "references/context/assignment-brief.png; report/sections/problem-data.tex",
                "Thời lượng: 35 giây.\nĐộng cơ turbofan gửi số đo cảm biến sau mỗi chu kỳ vận hành, gọi là cycle. Bảo trì "
                "quá sớm thì lãng phí, quá muộn thì động cơ có thể hỏng khi đang vận hành. Đại lượng cần ước lượng là RUL "
                f"(Remaining Useful Life), số cycle còn lại trước khi hỏng. Ví dụ động cơ train 1 hỏng ở cycle {failure}, "
                f"nên tại cycle {failure - 30} RUL bằng 30.\nKết quả chính trên tập test: hệ thống xếp "
                f"{ev['n_flagged_needs_maintenance']} động cơ cần bảo trì ngay và cả {ev['n_correct_needs_maintenance']} đều đúng.")
    plain_list(s, 0.92, 2.2, 6.7, 3.0, [
        f"Mỗi động cơ gửi {P['n_sensors']} số đo cảm biến sau mỗi chu kỳ vận hành (cycle).",
        "Bảo trì quá sớm gây lãng phí; bảo trì quá muộn có thể dẫn đến hỏng hóc khi vận hành.",
        "Cần ước lượng RUL (Remaining Useful Life): số cycle còn lại trước khi hỏng.",
    ], size=18)
    rect(s, 8.0, 2.2, 4.42, 2.8, TINT)
    text(s, 8.25, 2.4, 3.95, 2.5, [
        [("Ví dụ về RUL", {"bold": True, "color": BLUE})],
        f"Động cơ train 1 hỏng ở cycle {failure}.",
        f"Tại cycle {failure - 30}, RUL = {failure} − {failure - 30} = 30 cycle.",
    ], size=16, space_after=10)
    rect(s, 0.92, 5.35, 11.5, 1.2, BLUE)
    text(s, 1.2, 5.35, 11.0, 1.2, [
        [("Mục tiêu: ", {"bold": True, "color": GOLD}),
         ("xây dựng pipeline dữ liệu để phát hiện bất thường, dự đoán RUL và lập lịch bảo trì cho 100 động cơ test.",
          {"color": WHITE})],
    ], size=17, anchor=MSO_ANCHOR.MIDDLE)


# ============================================================================ 3. Quy trình
def slide_pipeline():
    s = content(1, 3, "Quy trình thực hiện gồm năm bước", "",
                "report/figures/pipeline.pdf; handoff/DATA_WALKTHROUGH.md",
                "Thời lượng: 30 giây.\nBước 1 tải dữ liệu từ NASA và kiểm tra checksum. Bước 2 lưu dữ liệu theo ba lớp "
                "Bronze, Silver, Gold dưới dạng Parquet và SQLite, đồng thời đưa lên HDFS. Bước 3 xử lý bằng Pandas, "
                "PySpark và MapReduce. Bước 4 phân tích: làm mượt EWMA, Health Index, mô hình HistGradientBoosting. Bước 5 "
                "tạo hai mức cảnh báo và lịch bảo trì.\nPhần 2 của bài (slide 4–12) trình bày bước 1–3; phần 3 (slide "
                "13–18) trình bày bước 4–5.")
    steps = [("1. Thu thập", "Tải từ NASA, kiểm tra MD5 và SHA-256"),
             ("2. Lưu trữ", "Bronze, Silver, Gold; Parquet, SQLite, HDFS"),
             ("3. Xử lý", "Pandas, PySpark, MapReduce trên YARN"),
             ("4. Phân tích", "EWMA, Health Index, mô hình RUL"),
             ("5. Kết quả", "Cảnh báo hai mức, lịch bảo trì")]
    w, gap = 2.05, 0.31
    for i, (head, body) in enumerate(steps):
        x = 0.92 + i * (w + gap)
        rect(s, x, 2.6, w, 1.9, BLUE if i == 4 else TINT)
        text(s, x + 0.15, 2.75, w - 0.3, 0.45, head, size=17, bold=True, color=WHITE if i == 4 else BLUE)
        text(s, x + 0.15, 3.3, w - 0.3, 1.1, body, size=14, color=WHITE if i == 4 else INK)
        if i:
            arrow(s, x - gap + 0.03, 3.55, x - 0.03, 3.55)
    for x0, x1, label in ((0.92, 0.92 + 3 * w + 2 * gap, "Dữ liệu và hạ tầng · slide 4–12"),
                          (0.92 + 3 * (w + gap), 12.42, "Phân tích và dự đoán · slide 13–18")):
        rect(s, x0, 4.75, x1 - x0, 0.08, GOLD, rounded=False)
        text(s, x0, 4.95, x1 - x0, 0.4, label, size=15, bold=True, color=GOLD_DARK, align=PP_ALIGN.CENTER)
    text(s, 0.92, 5.75, 11.5, 0.8, "Cùng một bộ dữ liệu được xử lý bằng Pandas, PySpark và MapReduce; kết quả được "
         "đối chiếu với nhau trước khi dùng cho phân tích.", size=15, color=SLATE)


# ============================================================================ 4. Nguồn dữ liệu
def slide_source():
    meta = {"FD001": ("1", "1 (HPC)"), "FD002": ("6", "1 (HPC)"), "FD003": ("1", "2 (HPC, Fan)"),
            "FD004": ("6", "2 (HPC, Fan)")}
    rows = [["Tập con", "Điều kiện vận hành", "Dạng lỗi", "Động cơ train", "Động cơ test", "Số bản ghi"]]
    for row in D["subsets"]:
        rows.append([row["name"], *meta[row["name"]], str(row["train_engines"]), str(row["test_engines"]),
                     vn(row["observations"], 0)])
    total = sum(row["observations"] for row in D["subsets"])
    rows.append(["Tổng", "", "", str(sum(r["train_engines"] for r in D["subsets"])),
                 str(sum(r["test_engines"] for r in D["subsets"])), vn(total, 0)])
    retrieved = DP["retrieved_at_utc"][:10]
    s = content(2, 4, "Nguồn dữ liệu: bộ NASA C-MAPSS gồm bốn tập con", "",
                "outputs/acquisition/manifest.json; outputs/cross_dataset/FD00x/etl/run_info.json",
                "Thời lượng: 40 giây.\nC-MAPSS là dữ liệu mô phỏng động cơ turbofan của NASA, có chuỗi cảm biến của từng "
                "động cơ từ lúc mới vận hành đến khi hỏng. Bốn tập con khác nhau về số điều kiện vận hành và số dạng lỗi. "
                f"Tổng cộng {vn(total, 0)} bản ghi.\nDữ liệu được tải bằng script từ {DP['source_url']} "
                f"(ngày {retrieved}, giờ UTC): gói {vn(DP['archive_bytes'] / 1e6, 1)} MB, {DP['n_files']} file; kiểm tra "
                "MD5 của gói và SHA-256 của từng file, ghi vào manifest.\nPhân tích chính dùng FD001; FD002–FD004 chỉ đi "
                "qua bước ETL.")
    table(s, rows, 0.92, 2.2, [1.4, 2.3, 2.1, 1.8, 1.8, 2.1], row_h=0.5, size=14, highlight=1, align_right_from=3)
    text(s, 0.92, 5.45, 11.5, 1.1, [
        f"Tải bằng script từ cổng dữ liệu mở của NASA: gói {vn(DP['archive_bytes'] / 1e6, 1)} MB, {DP['n_files']} file; "
        "kiểm tra MD5 của gói và SHA-256 của từng file.",
        "Phân tích chính dùng FD001 (dòng tô màu); ba tập còn lại chỉ đi qua bước ETL.",
    ], size=15, color=INK, space_after=6)


# ============================================================================ 5. Cấu trúc bản ghi
def slide_structure():
    sample = DP["raw_sample"]
    decimals = {"setting_1": 4, "setting_2": 4, "sensor_2": 2, "sensor_4": 2, "sensor_11": 2, "sensor_21": 4}
    rows = [sample["columns"] + ["…"]]
    for values in sample["rows"]:
        cells = [vn(v, decimals[c]) if c in decimals else str(int(v)) for c, v in zip(sample["columns"], values)]
        rows.append(cells + ["…"])
    n_train, n_obs = DP["shapes"]["train_targets"][0], DP["shapes"]["observations"][0]
    s = content(2, 5, f"Mỗi bản ghi là một cycle của một động cơ, gồm {sample['n_columns']} trường", "",
                "data/raw/train_FD001.txt; data/raw/readme.txt",
                "Thời lượng: 40 giây.\nBảng là bốn dòng đầu của file train_FD001.txt, đọc trực tiếp từ dữ liệu gốc. File "
                "text không có header, các trường cách nhau bằng khoảng trắng. Mỗi dòng gồm mã động cơ, số thứ tự cycle, "
                "3 thiết lập vận hành và 21 sensor; bảng chỉ hiện một phần các cột.\nMỗi tập con có ba file: train chạy "
                f"đến khi hỏng ({vn(n_train, 0)} dòng với FD001), test bị cắt trước khi hỏng ({vn(n_obs - n_train, 0)} dòng) "
                "và file RUL thật của test (100 dòng). File RUL chỉ được đọc sau khi đã lưu dự đoán.")
    table(s, rows, 0.92, 2.15, [1.2, 0.9, 1.45, 1.45, 1.3, 1.4, 1.35, 1.45, 0.5], row_h=0.42, size=13,
          align_right_from=0)
    text(s, 0.92, 4.4, 11.5, 0.4, "Bốn dòng đầu của train_FD001.txt; 26 trường: engine_id, cycle, setting_1–3, sensor_1–21.",
         size=13, color=SLATE)
    files = [("train_FD001.txt", f"{vn(n_train, 0)} dòng", "chạy đến khi hỏng"),
             ("test_FD001.txt", f"{vn(n_obs - n_train, 0)} dòng", "bị cắt trước khi hỏng"),
             ("RUL_FD001.txt", "100 dòng", "RUL thật của test, chỉ dùng để chấm")]
    for i, (name, size, role) in enumerate(files):
        x = 0.92 + i * 3.9
        rect(s, x, 5.0, 3.7, 1.5, TINT)
        text(s, x + 0.18, 5.1, 3.4, 0.4, name, size=15, bold=True, color=BLUE, font="Consolas")
        text(s, x + 0.18, 5.55, 3.4, 0.9, [size, role], size=15, space_after=2)


# ============================================================================ 6. Chất lượng
def slide_quality():
    q, E = DP["quality"], D["eda"]
    life = E["lifetime"]
    s = content(2, 6, f"Dữ liệu đầy đủ; tuổi thọ động cơ từ {vn(life['min'], 0)} đến {vn(life['max'], 0)} cycle", "",
                f"outputs/etl/silver/observations.parquet; {NB01}/engine_lifetimes.png",
                f"Thời lượng: 40 giây.\nKiểm tra trên {vn(q['rows'], 0)} bản ghi FD001: không có giá trị thiếu, không có "
                "khóa trùng, cycle của mỗi động cơ bắt đầu từ 1 và liên tục. Khóa của mọi bảng là (split, engine_id, "
                "cycle): cặp (engine_id, cycle) trùng giữa train và test nên phải thêm cột split.\nBảy cột không đổi ở cả "
                f"100 động cơ: sensor 1, 5, 10, 16, 18, 19 và setting_3; sensor 6 không đổi ở {E['partially_constant'][0]['constant_engines']}/100 động cơ. Các cột này "
                "không mang thông tin suy giảm nên không làm mượt; bảng đặc trưng vẫn giữ đủ 21 sensor.\n"
                "Hình do notebook 01 sinh "
                f"ra: tuổi thọ trung vị {vn(life['50%'], 0)} cycle, lệch phải. Nếu bảo trì cả đội ở cycle "
                f"{vn(life['50%'], 0)}, {E['n_fail_before_median']} động cơ đã hỏng trước đó, nên cần dự đoán riêng cho "
                "từng động cơ.")
    wide_figure(s, f"{NB01}/engine_lifetimes.png")
    numbers_below(s, [
        (str(q["missing_values"]), "giá trị thiếu", BLUE),
        (str(q["duplicate_keys"] + q["cycle_gaps"]), "khóa trùng hoặc cycle bị ngắt quãng", BLUE),
        (str(len(E["always_constant"])), "cột không đổi ở cả 100 động cơ, không mang thông tin suy giảm", GOLD_DARK),
    ])


# ============================================================================ 7. Kiến trúc lưu trữ
def slide_layers():
    sh = DP["shapes"]
    s = content(2, 7, "Dữ liệu được tổ chức theo ba lớp Bronze, Silver, Gold", "",
                "report/figures/data-flow.pdf; outputs/etl/; outputs/features/gold/rolling_features.parquet",
                "Thời lượng: 45 giây.\nSơ đồ là hình luồng dữ liệu trong báo cáo. Bronze giữ nguyên file text gốc, không "
                "ghi đè. Silver là bảng đã kiểm tra kiểu và khóa: train và test được gộp theo dòng (UNION) thành "
                f"{vn(sh['observations'][0], 0)} dòng, lưu Parquet. Gold là các bảng tổng hợp: engine_summary "
                f"{sh['engine_summary'][0]} dòng, mỗi động cơ một dòng, lưu cả Parquet và SQLite; bảng đặc trưng "
                f"{vn(sh['rolling_features'][0], 0)} dòng.\nParquet lưu theo cột nên đọc nhanh khi chỉ cần một vài sensor. "
                "Mỗi lần ghi đều qua thư mục tạm, kiểm tra xong mới thay bộ cũ, kèm run_info.json lưu hash dữ liệu đầu vào.")
    tall_figure(s, "slides/figures/data_flow.png")
    layers = [("Bronze", "File text gốc, giữ nguyên, không ghi đè"),
              ("Silver", f"Bảng đã kiểm tra: {vn(sh['observations'][0], 0)} × {sh['observations'][1]}, Parquet"),
              ("Gold", f"Bảng tổng hợp {sh['engine_summary'][0]} × {sh['engine_summary'][1]} (Parquet, SQLite); "
                       f"đặc trưng {vn(sh['rolling_features'][0], 0)} × {sh['rolling_features'][1]}")]
    for i, (name, body) in enumerate(layers):
        y = 2.2 + i * 1.45
        rect(s, 8.85, y, 3.57, 1.3, TINT)
        text(s, 9.05, y + 0.1, 3.2, 0.4, name, size=17, bold=True, color=BLUE)
        text(s, 9.05, y + 0.5, 3.25, 0.8, body, size=13)


# ============================================================================ 8. Hadoop
def slide_hadoop():
    mr, hd = DP["mapreduce"], HDFS
    paths = hd["hdfs_paths"]
    s = content(2, 8, f"MapReduce đọc dữ liệu từ HDFS, khớp Pandas {hd['mapreduce_vs_pandas']['rows_candidate']}/{hd['mapreduce_vs_pandas']['rows_candidate']} động cơ", "",
                "outputs/hdfs/summary.json; outputs/hdfs/mapreduce_job.txt; outputs/hdfs/hdfs_fsck.txt; "
                "outputs/hdfs/figures/yarn_applications.png",
                f"Thời lượng: 50 giây.\nCụm {hd['hadoop_version']} một node chạy bằng Docker: NameNode, DataNode cho HDFS; "
                "ResourceManager, NodeManager cho YARN. Hai file raw FD001 được đưa lên HDFS ở thư mục "
                f"{paths['bronze']}.\nJob MapReduce viết bằng Hadoop Streaming chạy trên YARN: {mr['map_tasks']} map task "
                f"và {mr['reduce_tasks']} reduce task, đọc {vn(mr['input_records'], 0)} bản ghi, ghi {mr['output_records']} "
                f"dòng tổng hợp, mất {vn(hd['seconds']['mapreduce_job'], 1)} giây. Spark đọc raw từ HDFS, ghi lớp Silver "
                f"dạng Parquet phân vùng theo split, mất {vn(hd['seconds']['spark_job'], 1)} giây.\nCả hai cho bảng "
                f"{hd['mapreduce_vs_pandas']['rows_candidate']} động cơ khớp tuyệt đối với Pandas. hdfs fsck: "
                f"{DP['fsck']['blocks']} block, trạng thái {DP['fsck']['status']}.\nGiới hạn: cụm một node nên không đo "
                "lợi ích của xử lý phân tán.")
    # Cắt bỏ logo và menu trái của giao diện YARN để chữ trong bảng ứng dụng đủ lớn.
    picture(s, "outputs/hdfs/figures/yarn_applications.png", 0.92, 2.05, 7.6, 3.15, crop=(205, 95, 1400, 578),
            align="left")
    code_box(s, 0.92, 5.3, 7.6, 1.35, [
        "$ hdfs dfs -ls -R /cmapss",
        f"{paths['bronze']}/train_FD001.txt",
        f"{paths['silver']}/split=train/part-00000-….snappy.parquet",
        f"$ hdfs fsck /cmapss   →   {DP['fsck']['blocks']} block, {DP['fsck']['status']}",
    ], size=11)
    numbers_right(s, [
        (f"{mr['map_tasks'] + mr['reduce_tasks']} task",
         f"{mr['map_tasks']} map task và {mr['reduce_tasks']} reduce task chạy trên YARN", BLUE),
        (f"{vn(mr['input_records'], 0)} → {mr['output_records']}", "bản ghi đầu vào → dòng tổng hợp đầu ra", BLUE),
        (f"{hd['mapreduce_vs_pandas']['rows_candidate']}/{hd['mapreduce_vs_pandas']['rows_reference']}",
         "động cơ khớp với kết quả Pandas", BLUE),
    ])


# ============================================================================ 9. PySpark
def slide_spark():
    parity = {row["name"]: row for row in P["parity"]}
    names = {"silver_observations": "observations (Silver)", "silver_train_targets": "train_targets (Silver)",
             "gold_engine_summary": "engine_summary (Gold)"}
    rows = [["Bảng", "Số dòng", "Sai khác lớn nhất", "Kết quả"]]
    for key, label in names.items():
        row = parity[key]
        rows.append([label, vn(row["rows"], 0), vn(row["max_abs_diff"], 0), "Khớp" if row["status"] == "passed" else "Lệch"])
    rows.append(["rolling_features (Gold)", vn(next(iter(ROLLING_ROWS)), 0), "< 10⁻⁷", "Khớp"])
    s = content(2, 9, "PySpark cho kết quả trùng khớp với Pandas", "",
                "outputs/etl/parity.json; outputs/benchmarks/20260929/run_info.json; src/cmapss/etl/rolling_spark.py",
                "Thời lượng: 45 giây.\nPipeline được viết hai lần với cùng logic: Pandas cho một máy và PySpark cho xử lý "
                "phân tán. Bộ đối chiếu so schema, khóa, số dòng và từng giá trị. Ba bảng ETL khớp tuyệt đối; bảng đặc "
                f"trưng {vn(next(iter(ROLLING_ROWS)), 0)} dòng khớp mọi ô với sai số cho phép 10⁻⁷.\nĐoạn code bên phải "
                "lấy từ rolling_spark.py: Window của Spark chia theo động cơ, sắp theo cycle và lấy 30 dòng gần nhất, "
                "tương đương rolling(30) của Pandas. Train và test được gộp bằng unionByName.")
    table(s, rows, 0.92, 2.2, [2.6, 1.15, 1.75, 1.0], row_h=0.5, size=13, align_right_from=1)
    code_box(s, 7.75, 2.2, 4.67, 3.35, [
        "window = (Window",
        "    .partitionBy(\"split\", \"engine_id\")",
        "    .orderBy(\"cycle\")",
        "    .rowsBetween(-29, 0))",
        "",
        "F.avg(sensor).over(window)",
        "F.stddev_samp(sensor).over(window)",
        "F.min(sensor).over(window)",
        "F.max(sensor).over(window)",
        "F.regr_slope(sensor, \"cycle\").over(window)",
    ], size=12)
    text(s, 7.75, 5.65, 4.67, 0.5, "Trích src/cmapss/etl/rolling_spark.py", size=12, color=SLATE)
    text(s, 0.92, 5.0, 6.5, 1.4, "Cùng một phép tính chạy bằng Pandas trên một máy và bằng PySpark, cho kết quả như nhau; "
         "vì vậy có thể chuyển sang Spark khi dữ liệu lớn hơn mà không đổi kết quả phân tích.", size=15, color=INK)


# ============================================================================ 10. Truy vấn
def slide_query():
    sql_lines = DP["sql"]["query"].splitlines()
    result = [["engine_id", "n_cycles", "cycle", "rul"]] + [[str(v) for v in row] for row in DP["sql"]["rows"]]
    mean = DP["rolling_mean_sensor_11"]
    s = content(2, 10, "SQL và Spark SQL trả về đúng nhãn RUL và đặc trưng window", "",
                "outputs/etl/gold/etl.sqlite; outputs/etl/silver/observations.parquet; report/sections/data-processing.tex",
                "Thời lượng: 45 giây.\nTruy vấn bên trái chạy trên file etl.sqlite của lớp Gold: JOIN bảng tổng hợp động "
                "cơ với bảng nhãn theo khóa split và engine_id, lấy RUL của động cơ train 1 tại cycle 1, 100 và 192. Kết "
                "quả dưới truy vấn là kết quả chạy thật.\nBên phải là Spark SQL trên file Parquet, dùng window function để "
                "tính trung bình 30 cycle gần nhất của sensor 11, chính là đặc trưng sensor_11__mean. Điều kiện lọc cycle "
                "phải đặt ở truy vấn ngoài; nếu đặt cùng cấp với window function thì window chỉ thấy các dòng đã lọc và "
                "cho kết quả sai.")
    text(s, 0.92, 2.1, 5.6, 0.35, "SQL trên SQLite (lớp Gold)", size=15, bold=True, color=BLUE)
    code_box(s, 0.92, 2.5, 5.6, 2.25, sql_lines, size=11)
    table(s, result, 0.92, 4.9, [1.4, 1.4, 1.2, 1.0], row_h=0.36, size=12, align_right_from=0)
    text(s, 6.82, 2.1, 5.6, 0.35, "Spark SQL trên Parquet (lớp Silver)", size=15, bold=True, color=BLUE)
    code_box(s, 6.82, 2.5, 5.6, 2.9, [
        "SELECT cycle, mean30 FROM (",
        "  SELECT cycle, AVG(sensor_11) OVER w AS mean30",
        "  FROM observations",
        "  WHERE split = 'train' AND engine_id = 1",
        "  WINDOW w AS (PARTITION BY split, engine_id",
        "    ORDER BY cycle",
        "    ROWS BETWEEN 29 PRECEDING AND CURRENT ROW))",
        "WHERE cycle IN (30, 100);",
    ], size=11)
    text(s, 6.82, 5.55, 5.6, 0.9, [
        f"Kết quả: cycle 30 → {vn(mean['30'], 3)}; cycle 100 → {vn(mean['100'], 3)}",
        "Trùng với đặc trưng sensor_11__mean do Pandas tính.",
    ], size=14, space_after=4)


# ============================================================================ 11. Biến đổi
def slide_transform():
    sh, fr = DP["shapes"], P["feature_rows"]
    n_train = sh["train_targets"][0]

    def size(key):
        return f"{vn(sh[key][0], 0)} × {sh[key][1]}"

    life1 = M["example_warning"]["failure_cycle"]  # động cơ train 1
    n_rows1 = life1 - (P["window"] - 1)
    rows = [["#", "Phép biến đổi", "Kết quả", "Kích thước"],
            ["1", "Đọc 2 file text, gán tên và kiểu 26 trường, thêm cột split", "2 bảng có kiểu",
             f"{vn(n_train, 0)} và {vn(sh['observations'][0] - n_train, 0)} dòng"],
            ["2", "Gộp train và test theo dòng (UNION), sắp theo khóa", "observations", size("observations")],
            ["3", "GROUP BY động cơ; RUL = cycle cuối − cycle hiện tại", "train_targets", size("train_targets")],
            ["4", "GROUP BY (split, engine_id): số cycle, cycle đầu, cycle cuối", "engine_summary", size("engine_summary")],
            ["5", f"Window {P['window']} cycle: {P['n_statistics']} thống kê × {P['n_sensors']} sensor",
             "rolling_features", size("rolling_features")],
            ["6", "JOIN đặc trưng train với nhãn theo đủ 3 khóa, một-một", "X, y",
             f"{vn(fr['rolling_features_train'], 0)} × {P['feature_columns']}"]]
    s = content(2, 11, f"Sáu bước biến đổi từ file text thành {vn(fr['rolling_features_train'], 0)} mẫu huấn luyện", "",
                "report/sections/data-processing.tex; outputs/etl/run_info.json; outputs/features/run_info.json",
                "Thời lượng: 45 giây.\nĐây là bước T trong ETL. Bước 1–5 được viết bằng cả Pandas và PySpark. Ở bước 5, "
                f"mỗi dòng đặc trưng là 6 thống kê (giá trị cuối, trung bình, độ lệch chuẩn, min, max, độ dốc) của "
                f"{P['n_sensors']} sensor trên {P['window']} cycle gần nhất, chỉ dùng quá khứ. 29 cycle đầu mỗi động cơ "
                f"chưa đủ window nên không tạo dòng: {vn(n_train, 0)} − 100 × 29 = {vn(fr['rolling_features_train'], 0)} "
                "dòng train.\nBước 6 bắt buộc quan hệ một-một. Nếu chỉ join theo engine_id, động cơ train 1 sẽ thành "
                f"{n_rows1} × {life1} = {vn(n_rows1 * life1, 0)} dòng thay vì {n_rows1} dòng.")
    table(s, rows, 0.92, 2.15, [0.5, 6.3, 2.2, 2.5], row_h=0.58, size=14, align_right_from=3)


# ============================================================================ 12. Mở rộng quy mô
def slide_scale():
    summary = D["benchmark_scaling"]["summary"]
    common, ratio, stop = summary["largest_common_scale"], summary["pandas_over_spark_ratio"], summary["first_pandas_stop_scale"]
    assert common is not None and stop is not None, "benchmark chưa đủ lượt chạy"
    spark_stop = summary["spark_seconds_at_stop"]
    s = content(2, 12, f"Khi dữ liệu tăng {common} lần, Spark tính nhanh hơn Pandas {vn(ratio, 1)} lần", "",
                "slides/figures/benchmark_scaling.png; outputs/benchmarks/20261021/compute_only.csv",
                "Thời lượng: 50 giây.\nDữ liệu FD001 được nhân bản 1, 10, 50, 100 và 200 lần, rồi cho Pandas và Spark cùng "
                "đọc, tổng hợp và tạo đặc trưng. Thời gian chỉ gồm phần tính toán, không gồm ghi file. Dấu x là lượt quá "
                f"giới hạn 30 phút.\nỞ {common}×, Pandas chậm hơn Spark {vn(ratio, 1)} lần. Ở {stop}×, Pandas không xong "
                f"trong 30 phút; Spark xong sau {vn(spark_stop, 0)} giây.\nGiới hạn: Pandas chạy 1 luồng, Spark local[4] "
                "chạy 4 luồng trên cùng một máy, nên một phần chênh lệch đến từ song song hóa. Pandas tính slope bằng một "
                "hàm Python cho từng window, bước này chiếm phần lớn thời gian; Spark dùng hàm dựng sẵn regr_slope. Vì "
                "vậy hệ số lớn hơn số luồng. Mỗi mức chỉ đo một lượt. Với 100 động cơ hiện tại, "
                "Pandas vẫn đủ nhanh và lịch bảo trì không đổi; Spark cần khi số động cơ tăng hàng chục lần.")
    wide_figure(s, "slides/figures/benchmark_scaling.png")
    numbers_below(s, [
        (f"{vn(ratio, 1)} lần", f"Spark nhanh hơn Pandas ở {common}× dữ liệu", BLUE),
        ("> 30 phút", f"Pandas ở {stop}×; Spark xong sau {vn(spark_stop, 0)} giây", RED),
        ("Như nhau", "lịch bảo trì với 100 động cơ hiện tại khi dùng Pandas", SLATE),
    ])


# ============================================================================ 13. Làm mượt
def slide_denoise():
    reductions = [float(r["reduction_pct"]) for r in NOISE.values()]
    s4 = NOISE["sensor_4"]
    s = content(3, 13, f"Làm mượt EWMA giảm khoảng {vn(DEG['noise_reduction_pct_median'], 0)}% dao động ngắn hạn", "",
                "outputs/degradation/figures/denoise_before_after.png; outputs/degradation/tables/noise_summary.csv",
                "Thời lượng: 35 giây.\nHình do module degradation sinh ra: sensor 11 và sensor 4 của động cơ train 1. "
                "Đường xám là tín hiệu gốc, đường xanh sau EWMA span 10. Công thức s_t = α·x_t + (1 − α)·s_{t−1}, "
                "α = 2/(span + 1) ≈ 0,18; chỉ dùng giá trị hiện tại và quá khứ.\nĐo bằng độ lệch chuẩn của chênh lệch "
                f"giữa hai cycle liền kề. Ví dụ sensor 4: {vn(float(s4['diff_std_raw']))} xuống "
                f"{vn(float(s4['diff_std_smoothed']))}. Trên 14 sensor, mức giảm từ {vn(min(reductions), 1)}% đến "
                f"{vn(max(reductions), 1)}%. Với dao động độc lập, span 10 cho mức giảm lý thuyết khoảng 86,5%, nên con "
                "số này kiểm chứng tham số làm mượt, không phải phát hiện riêng về dữ liệu.\nKhông làm mượt thì cảnh báo mức một đến muộn hơn: trung vị "
                f"{vn(AN_RAW['lead_median'], 1)} cycle trước khi hỏng thay vì {vn(AN['lead_median'], 0)} cycle.")
    wide_figure(s, "outputs/degradation/figures/denoise_before_after.png")
    numbers_below(s, [
        (f"−{vn(DEG['noise_reduction_pct_median'], 0)}%", "dao động ngắn hạn sau khi làm mượt (trung vị 14 sensor)", BLUE),
        (f"{vn(AN['lead_median'], 0)} cycle", "cảnh báo mức một báo trước khi hỏng, có làm mượt", BLUE),
        (f"{vn(AN_RAW['lead_median'], 1)} cycle", "nếu không làm mượt", RED),
    ])


# ============================================================================ 14. Xu hướng
def slide_trend():
    st = D["statistics"]
    s = content(3, 14, f"{len(SAME_DIRECTION)} sensor thay đổi cùng chiều ở cả 100 động cơ khi gần hỏng", "",
                f"{NB01}/selected_sensor_trajectories.png; outputs/statistics/tables/sensor_inference.csv",
                "Thời lượng: 30 giây.\nHình do notebook 01 sinh ra: tín hiệu gốc của ba động cơ train dài 178, 199 và 229 "
                "cycle. Sensor 11 tăng dần và sensor 12 giảm dần khi động cơ gần hỏng.\nKiểm định: với mỗi động cơ, lấy "
                f"trung bình {st['comparison_cycles']} cycle cuối trừ {st['comparison_cycles']} cycle đầu; bootstrap "
                f"{vn(st['resamples'], 0)} lần cho khoảng tin cậy; kiểm định Wilcoxon; "
                f"hiệu chỉnh Benjamini–Hochberg với q = {vn(st['fdr_q'], 2)} trên {st['n_sensors']} sensor. Kết quả: "
                f"{st['n_reject']} sensor thay đổi có "
                f"ý nghĩa; {len(SAME_DIRECTION)} sensor thay đổi cùng chiều ở cả 100 động cơ; {st['n_degenerate']} sensor "
                "không đổi nên không "
                "kiểm định được.")
    wide_figure(s, f"{NB01}/selected_sensor_trajectories.png", crop=(0, 0, 1817, 535))
    numbers_below(s, [
        (str(len(SAME_DIRECTION)), "sensor thay đổi cùng chiều ở cả 100 động cơ", BLUE),
        (f"{st['n_reject']}/{st['n_sensors']}", f"sensor thay đổi có ý nghĩa thống kê (q = {vn(st['fdr_q'], 2)})", BLUE),
    ])


# ============================================================================ 15. Bất thường
def slide_anomaly():
    assert AN["n_detected"] == AN["n_engines"]
    anomaly_cycle = int(next(row for row in D["example_anomaly"] if row["engine_id"] == 1)["detection_cycle"])
    failure = int(M["example_warning"]["failure_cycle"])
    s = content(3, 15, f"Health Index phát hiện bất thường ở cả {AN['n_engines']} động cơ train", "",
                "outputs/degradation/figures/anomaly_health_index.png; outputs/degradation/tables/anomaly_events.csv",
                "Thời lượng: 45 giây.\nHình do module degradation sinh ra. Health Index gộp 21 sensor bằng StandardScaler "
                "rồi PCA một thành phần; chỉ số giảm khi động cơ suy giảm. Không dùng nhãn RUL.\nQuy tắc cố định trước: "
                f"mức tham chiếu là {ACFG['baseline_cycles']} cycle đầu của chính động cơ đó. Khi chỉ số giảm quá "
                f"{vn(ACFG['z_threshold'], 0)} độ lệch chuẩn và giữ {ACFG['persistence']} cycle liên tiếp thì phát cảnh "
                f"báo mức một.\nHình trái, động cơ train 1: phát hiện ở cycle {anomaly_cycle}, hỏng ở cycle {failure}. "
                f"Hình phải: trung vị {vn(AN['lead_median'], 0)}, nhỏ nhất {vn(AN['lead_min'], 0)}, lớn nhất "
                f"{vn(AN['lead_max'], 0)} cycle.\nCả chín cấu hình trong lưới (2/3/4σ × 3/5/10 cycle) đều phát hiện đủ 100 động cơ. Giới hạn: chưa "
                "đo được báo động giả vì dữ liệu không ghi thời điểm bắt đầu suy giảm.")
    wide_figure(s, "outputs/degradation/figures/anomaly_health_index.png")
    numbers_below(s, [
        (f"{AN['n_detected']}/{AN['n_engines']}", "động cơ train được phát hiện trước khi hỏng", BLUE),
        (f"{vn(AN['lead_median'], 0)} cycle", "thời gian báo trước (trung vị)", BLUE),
        (f"cycle {anomaly_cycle}", f"động cơ train 1 được phát hiện; hỏng ở cycle {failure}", GOLD_DARK),
    ])


# ============================================================================ 16. RUL
def slide_rul():
    f = M["final"]
    cv = {r["model_name"]: r for r in M["cv"]}
    s = content(3, 16, f"Mô hình HistGradientBoosting dự đoán RUL với MAE {vn(f['mean_engine_mae'])} cycle", "",
                "outputs/modeling/figures/final_true_vs_predicted.png; outputs/modeling/tables/final_metrics.json",
                "Thời lượng: 45 giây.\nHình do module modeling sinh ra: mỗi điểm là một động cơ test, dự đoán tại cycle "
                "cuối; điểm càng gần đường chéo càng đúng.\nMô hình được chọn bằng cross-validation 5 fold chia theo "
                f"động cơ (GroupKFold): HGB có MAE {vn(cv['hist_gradient_boosting']['mean_engine_mae'])} cycle, thấp hơn "
                f"Ridge {vn(M['ridge_minus_hgb_mae'])} cycle; baseline trung vị {vn(cv['median_baseline']['mean_engine_mae'])}.\n"
                f"Trên test: MAE {vn(f['mean_engine_mae'])} cycle, khoảng tin cậy 95% [{vn(M['ci_mae']['ci_lower'])}; "
                f"{vn(M['ci_mae']['ci_upper'])}], RMSE {vn(f['rmse'])}. Với {M['low_rul']['n_engines']} động cơ có RUL thật "
                f"≤ 30, MAE {vn(M['low_rul']['mae'])} cycle. MAE test thấp hơn MAE cross-validation vì test chỉ có một "
                "điểm cuối mỗi chuỗi.")
    tall_figure(s, "outputs/modeling/figures/final_true_vs_predicted.png")
    numbers_right(s, [
        (f"{vn(f['mean_engine_mae'])} cycle", "MAE trên 100 động cơ test", BLUE),
        (f"{vn(M['low_rul']['mae'])} cycle", f"MAE với {M['low_rul']['n_engines']} động cơ có RUL thật ≤ 30", RED),
        (f"{vn(cv['hist_gradient_boosting']['mean_engine_mae'])} cycle", "MAE cross-validation 5 fold theo động cơ", SLATE),
    ])


# ============================================================================ 17. Cảnh báo sớm
def slide_warning():
    tm = M["warning"]["timing"]
    w30, w40 = WARN["hgb126_30"], WARN["hgb126_40"]
    example = M["example_warning"]
    s = content(3, 17, f"Cảnh báo sớm: ngưỡng 40 cycle nâng tỷ lệ báo kịp thời lên {vn(w40['on_time_rate'] * 100, 0)}%", "",
                "outputs/modeling/figures/warning_trajectory.png; outputs/course_alignment/20261007/warning_fixed_policies.csv",
                "Thời lượng: 45 giây.\nHình do module modeling sinh ra: biểu đồ xu hướng cảnh báo sớm của động cơ train 1, "
                f"dự đoán OOF. Cảnh báo mức hai phát khi RUL dự đoán ≤ 30; động cơ train 1 được cảnh báo ở cycle "
                f"{example['first_warning_cycle']}, hỏng ở {example['failure_cycle']}.\nTrên 100 động cơ train ở ngưỡng 30: "
                f"{tm['before_threshold']} báo trước, {tm['at_threshold']} đúng lúc, {tm['after_threshold']} muộn nhưng vẫn "
                f"trước khi hỏng. Cảnh báo muộn rơi nhiều vào các động cơ hỏng sớm; báo cáo nêu đây là giả thuyết, "
                "chưa chứng minh nguyên nhân.\n"
                f"Ngưỡng 40 nâng tỷ lệ kịp thời từ {vn(w30['on_time_rate'] * 100)}% lên {vn(w40['on_time_rate'] * 100)}%, "
                f"đổi lại tỷ lệ cycle báo quá sớm từ {vn(w30['early_alarm_burden'] * 100)}% lên "
                f"{vn(w40['early_alarm_burden'] * 100)}%. Chọn ngưỡng thực tế cần chi phí bảo trì.")
    tall_figure(s, "outputs/modeling/figures/warning_trajectory.png")
    numbers_right(s, [
        (f"{tm['after_threshold']}/100", "động cơ train bị cảnh báo muộn ở ngưỡng 30", RED),
        (f"{vn(w30['on_time_rate'] * 100, 0)}% → {vn(w40['on_time_rate'] * 100, 0)}%", "tỷ lệ cảnh báo kịp thời khi nâng ngưỡng lên 40", BLUE),
        (f"{vn(w30['early_alarm_burden'] * 100, 1)}% → {vn(w40['early_alarm_burden'] * 100, 1)}%", "tỷ lệ cycle báo quá sớm", SLATE),
    ])


# ============================================================================ 18. Lịch bảo trì
def slide_maintenance():
    ev, counts, plan = MNT["evaluation"], MNT["status_counts"], MNT["planning"]
    missed, late = MNT["missed"], MNT["late"]
    s = content(3, 18, f"{ev['n_correct_needs_maintenance']} chỉ định bảo trì ngay đều đúng, còn sót "
                f"{missed['n_engines']}/{ev['n_actual_needs_maintenance']} động cơ", "",
                "outputs/maintenance/figures/fleet_status.png; outputs/maintenance/tables/maintenance_schedule.csv",
                "Thời lượng: 50 giây; sau slide này chuyển sang demo dashboard khoảng 2 phút.\nHình do module maintenance "
                "sinh ra: 100 động cơ test xếp theo RUL dự đoán. Cần bảo trì nếu RUL dự đoán ≤ 30; theo dõi tăng cường nếu "
                f"Health Index đã báo bất thường; còn lại bình thường. Kết quả: {counts['Cần bảo trì']} cần bảo trì, "
                f"{counts['Theo dõi tăng cường']} theo dõi, {counts['Bình thường']} bình thường; thêm {plan['1–25']} động "
                f"cơ đến hạn trong 25 cycle tới.\nĐối chiếu với nhãn test sau khi đã lưu lịch: "
                f"{ev['n_actual_needs_maintenance']} động cơ thật sự cần bảo trì; {ev['n_correct_needs_maintenance']} được "
                f"báo đúng; {missed['n_engines']} động cơ bị sót có RUL dự đoán {vn(missed['rul_pred_min'])}–"
                f"{vn(missed['rul_pred_max'])}. Ngưỡng 40 (chọn từ train) bắt đủ {T40['correct']}/{T40['actual']}, báo sớm "
                f"{T40['early']}.\nGiới hạn: {late['n_engines']} động cơ còn xa mới hỏng có lịch muộn hơn lúc hỏng "
                f"{late['delay_min']}–{late['delay_max']} cycle; lịch được tính lại khi có dữ liệu mới.")
    wide_figure(s, "outputs/maintenance/figures/fleet_status.png")
    numbers_below(s, [
        (f"{ev['n_correct_needs_maintenance']}/{ev['n_flagged_needs_maintenance']}", "động cơ được xếp bảo trì ngay đều đúng", BLUE),
        (str(missed["n_engines"]), f"động cơ bị sót, RUL dự đoán {vn(missed['rul_pred_min'], 0)}–"
         f"{vn(missed['rul_pred_max'], 0)} cycle", GOLD_DARK),
        (f"{T40['correct']}/{T40['actual']}", "bắt đủ nếu dùng ngưỡng 40 cycle", BLUE),
    ])


# ============================================================================ 19. Kết luận
def slide_conclusion():
    f, ev = M["final"], MNT["evaluation"]
    s = content(4, 19, "Đạt mục tiêu trên FD001; còn hạn chế về chi phí và quy mô", "",
                "report/sections/conclusion.tex; report/sections/discussion.tex",
                "Thời lượng: 35 giây.\nKết luận: pipeline dữ liệu chạy được trên Pandas, PySpark và Hadoop với kết quả "
                "khớp; Health Index phát hiện bất thường sớm; mô hình RUL và lịch bảo trì cho kết quả tốt ở vùng gần hỏng.\n"
                "Hạn chế: dữ liệu mô phỏng và không có chi phí bảo trì; mô hình thường dự đoán cao hơn thực tế; cụm Hadoop "
                "chỉ có một node.\nHướng phát triển: chọn ngưỡng theo chi phí, chạy trên cụm nhiều node, nhận dữ liệu "
                "streaming.")
    columns = [
        ("Kết luận", [
            "Pipeline chạy trên Pandas, PySpark, Hadoop; bảng ETL khớp tuyệt đối, bảng đặc trưng sai số < 10⁻⁷.",
            f"Health Index phát hiện bất thường ở {AN['n_detected']}/{AN['n_engines']} động cơ, trung vị "
            f"{vn(AN['lead_median'], 0)} cycle trước khi hỏng.",
            f"RUL: MAE {vn(f['mean_engine_mae'])} cycle; {ev['n_correct_needs_maintenance']} chỉ định bảo trì ngay "
            f"đều đúng, sót {MNT['missed']['n_engines']}/{ev['n_actual_needs_maintenance']} động cơ.",
        ]),
        ("Hạn chế", [
            "Dữ liệu mô phỏng, chưa có chi phí bảo trì.",
            f"Mô hình thường dự đoán RUL cao hơn thực tế ({M['n_over']}/100 động cơ test).",
            "Hadoop và Spark mới chạy trên một máy.",
        ]),
        ("Hướng phát triển", [
            "Chọn ngưỡng cảnh báo theo chi phí bảo trì.",
            "Triển khai trên cụm nhiều node.",
            "Nhận dữ liệu cảm biến dạng streaming.",
        ]),
    ]
    for i, (head, items) in enumerate(columns):
        x = 0.92 + i * 3.9
        rect(s, x, 2.15, 3.7, 4.4, BLUE if i == 0 else TINT)
        fg = WHITE if i == 0 else INK
        text(s, x + 0.2, 2.3, 3.3, 0.45, head, size=19, bold=True, color=GOLD if i == 0 else BLUE)
        text(s, x + 0.2, 2.9, 3.3, 3.5, [f"–  {item}" for item in items], size=15, color=fg, space_after=12)


# ============================================================================ 20. Kết thúc
def slide_end():
    slide = PRS.slides.add_slide(LAYOUT["Slide End"])
    title = slide.placeholders[0]
    title.left, title.top, title.width, title.height = Inches(0.69), Inches(2.0), Inches(6.2), Inches(2.2)
    title.text = "Trân trọng cảm ơn quý thầy cô đã lắng nghe"
    p = title.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    p.runs[0].font.size, p.runs[0].font.bold, p.runs[0].font.name = Pt(36), True, FONT
    p.runs[0].font.color.rgb = rgb(BLUE)
    text(slide, 0.69, 4.3, 6.2, 0.5, "Kính mời quý thầy cô đặt câu hỏi", size=20, color=INK)
    text(slide, 0.69, 5.2, 6.2, 0.9, [
        [("Giảng viên hướng dẫn: ", {"color": SLATE}), ("TS. Võ Thị Thanh Vân", {"bold": True, "color": BLUE})],
        [("Học viên: ", {"color": SLATE}), ("Phan Nguyễn Mai Phương · MSHV 25002711", {"color": INK})],
    ], size=13, space_after=4)
    slide.notes_slide.notes_text_frame.text = (
        "Thời lượng: 10 giây.\nCảm ơn quý thầy cô, mời đặt câu hỏi. Không chạy lại huấn luyện khi trình bày.")


def main() -> None:
    for build in (slide_cover, slide_problem, slide_pipeline, slide_source, slide_structure, slide_quality,
                  slide_layers, slide_hadoop, slide_spark, slide_query, slide_transform, slide_scale,
                  slide_denoise, slide_trend, slide_anomaly, slide_rul, slide_warning, slide_maintenance,
                  slide_conclusion, slide_end):
        build()
    assert len(PRS.slides) == TOTAL, len(PRS.slides)
    PRS.core_properties.author = "Phan Nguyễn Mai Phương"
    PRS.core_properties.title = "Đề tài 8 · Phân tích dữ liệu IoT và bảo trì dự đoán"
    PRS.save(str(OUT))
    print(f"wrote {OUT.relative_to(ROOT)} ({len(PRS.slides)} slides)")


if __name__ == "__main__":
    main()
